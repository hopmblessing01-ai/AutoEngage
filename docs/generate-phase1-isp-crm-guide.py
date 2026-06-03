"""Generate balanced Phase 1 ISP CRM job implementation guide (DOCX)."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

OUT = r"d:\RMH\Upwork Project\Blessing\automation-business\docs\Phase-1-ISP-CRM-Implementation-Guide.docx"


def build():
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)

    title = doc.add_heading("How to Deliver Phase 1: ISP CRM Automation", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "A practical work guide for completing the Upwork job: connect ISP Excel, Google Sheets CRM, "
        "RingCentral, and Zapier. Phase 1 only — do not build Smith.ai, Google Calendar, or SMS unless the client opens a new phase."
    )

    def h(text, level=1):
        doc.add_heading(text, level=level)

    def p(text):
        doc.add_paragraph(text)

    def bullets(items):
        for item in items:
            doc.add_paragraph(item, style="List Bullet")

    def steps(items):
        for item in items:
            doc.add_paragraph(item, style="List Number")

    # --- INTRO ---
    h("What you are building")
    p(
        "The client already has a Google Sheets CRM with tabs, scripts, and color coding. Your job is not to rebuild it. "
        "You are wiring four pieces together: (1) their ISP Excel customer export, (2) the CRM as the working system for agents, "
        "(3) RingCentral call activity, and (4) Zapier as the bridge from RingCentral into Sheets. Agents must look up customers on incoming calls, "
        "log notes, update call and retention status, and avoid duplicate records."
    )

    h("Golden rules")
    bullets([
        "Backup the CRM (File → Make a copy) before any script or structure change.",
        "Audit existing tabs and Apps Script first; extend rather than replace unless the client approves in writing.",
        "Any removal needs: why, replacement, how it is better, how it scales — then wait for approval.",
        "Get a signed column mapping (Excel → CRM) before coding the import.",
        "Test with real sample rows and real test calls before calling Phase 1 done.",
    ])

    doc.add_page_break()

    # --- WEEK BY WEEK HOW TO DO THE JOB ---
    h("How to run the project (step by step)")

    h("Step 1 — Kickoff and access (Days 1–2)", level=2)
    p("Request these before deep work:")
    bullets([
        "Editor access to the live CRM; create your own sandbox copy to build in first.",
        "A recent ISP Excel export (anonymized is fine) and how often they update it.",
        "RingCentral admin or integration access; list of extensions and agents.",
        "Zapier login and plan limits (tasks/month, number of Zaps).",
        "Their canonical customer ID (account number vs customer number vs phone).",
        "Lists they use (or want) for call disposition and retention/customer status.",
        "A short description of what agents do today when a call comes in.",
    ])
    p("Deliverable: confirm scope in writing — Phase 1 is Excel import, CRM upsert, duplicates, RingCentral logging, agent workflow. No Smith.ai or Calendar.")

    h("Step 2 — Audit the CRM (Days 2–4)", level=2)
    p("Open every tab and document:")
    steps([
        "Tab name, purpose, who updates it, approximate row count.",
        "Which tab is the master customer list (one row per customer).",
        "Existing import tabs and any import scripts or menus.",
        "All Apps Script: open Extensions → Apps Script, list triggers (onEdit, time-driven, menu).",
        "Conditional formatting and protected ranges (often tied to status — do not break blindly).",
        "How agents search for customers today.",
    ])
    p("Output a short audit (1–2 pages): current flow, risks, what you will add vs leave untouched. Client signs off before build.")

    h("Step 3 — Audit the Excel file (Days 3–5)", level=2)
    steps([
        "Identify the correct export sheet and required columns.",
        "Check account numbers (Excel may strip leading zeros — force text format).",
        "Check phone columns (multiple phones? inconsistent formats?).",
        "Check address and name fields; note blank rates.",
        "Map which Excel columns should update which CRM columns on each import.",
    ])
    p("Build a mapping table: Excel column → CRM column → rule (trim, date, status map). Client approves this document — it is your contract for import logic.")

    h("Step 4 — Design the data model (Days 5–7)", level=2)
    p("Use the client’s tab names when they already exist. If you add tabs, keep names clear:")
    bullets([
        "Master customers — canonical record; upsert target.",
        "Import_Staging — paste or CSV load from Excel; cleared or archived after each run.",
        "Import_Log — batch date, rows processed, created, updated, errors.",
        "Call_Log — one row per RingCentral call (call ID, time, phone, agent, duration, disposition, link to customer).",
        "Notes — timestamped notes linked to customer (preferred over one giant notes cell).",
        "Settings — status dropdown values, extension-to-agent map, match rules.",
        "Duplicate_Review — rows where import or call matching found 2+ possible customers.",
        "Unmatched calls — optional filter/view for calls with no customer link yet.",
    ])
    p("Add only the tabs/columns you need; do not rename or delete client tabs without approval.")

    h("Step 5 — Build the Excel import (Days 7–12)", level=2)
    p("Recommended for Phase 1: manual, reliable process.")
    steps([
        "Agent or admin exports Excel to CSV (or copies into Import_Staging).",
        "Custom menu in Sheets: Run Import.",
        "Script assigns a batch ID, normalizes phones and account numbers, validates required fields.",
        "For each staging row, search master using match order (see below).",
        "0 matches → create new customer (if client allows auto-create).",
        "1 match → update only the columns agreed in the mapping (do not overwrite immutable IDs).",
        "2+ matches → write row to Duplicate_Review; do not change master until a supervisor resolves.",
        "Write summary to Import_Log; show errors (bad phone, missing account #, etc.).",
    ])
    p("Apps Script to write: normalizePhone(), normalizeAccountId(), findMatches(), upsertFromStaging(), and LockService so two people cannot run import at once.")

    h("Step 6 — Duplicate prevention (Days 10–14)", level=2)
    p("Match priority (confirm with client, implement in this order):")
    steps([
        "Exact account/customer number (normalized text).",
        "Exact primary phone (normalized 10-digit or E.164 helper column).",
        "Secondary phone if they use one.",
        "Name + address together only if the above fail — avoid fuzzy auto-merge in Phase 1.",
    ])
    p("Prevention: never insert a second row with the same account number or same primary phone. Add menu item Duplicate Scan for supervisors. Do not delete duplicates without approval — merge or mark inactive.")

    h("Step 7 — RingCentral + Zapier (Days 12–18)", level=2)
    p("Keep Zaps simple to stay within task limits:")
    bullets([
        "Zap A — Trigger: RingCentral new call (inbound and/or missed). Action: Create row on Call_Log with call ID, datetime, from/to number, extension, duration.",
        "Zap B (if needed) — Lookup customer in Sheets by phone; update Call_Log with customer ID, or update last_call date on master.",
    ])
    p("If Zapier cannot match reliably (multiple tabs, fuzzy logic), add a small Apps Script web app or a post-import Script that runs on new Call_Log rows and sets customer_id.")
    p("Test with the client on real extensions: known customer phone, unknown phone, missed call. Confirm rows appear within a minute.")

    h("Step 8 — Status, notes, and agent workflow (Days 15–20)", level=2)
    p("On Settings tab, list allowed values. Use Data validation on CRM entry columns.")
    bullets([
        "Call disposition examples: Answered, Missed, Voicemail, Callback scheduled, Wrong number.",
        "Retention status examples: Active, At risk, Service issue, Cancel requested, Retained, Churned.",
    ])
    p("Agent workflow you implement and document:")
    steps([
        "Call arrives — agent answers in RingCentral.",
        "Open CRM bookmark → use menu Find by Phone (caller ID) or filter on phone/account column.",
        "Customer found — review last note and retention status; talk; then add note (Notes tab or form), set disposition and status if changed.",
        "Customer not found — follow client rule: create provisional record or log as unmatched and link after call.",
        "Follow-up — set follow-up date/flag; supervisor uses a filtered view for due items.",
    ])
    p("Deliver a 1–2 page Agent Quick Guide and a short Admin Guide (how to run import, read Import_Log, fix duplicates, reconnect Zapier).")

    h("Step 9 — Test and handoff (Days 18–22)", level=2)
    p("Run these tests and log results on a Test_Log tab:")
    bullets([
        "Import brand-new customer from Excel.",
        "Import update to phone or address on existing account — no duplicate row.",
        "Staging row that matches two customers — lands in Duplicate_Review only.",
        "Inbound call from known number — Call_Log linked to customer.",
        "Inbound call from unknown number — logged, agent can attach.",
        "Agent adds note and changes retention status.",
    ])
    p("Handoff: live CRM (or promoted sandbox), working Zaps, Scripts documented, SOPs, Test_Log, and a short list of Phase 2 ideas (Smith.ai, Calendar) without building them.")

    doc.add_page_break()

    # --- TECHNICAL REFERENCE (medium detail) ---
    h("Technical reference (while you build)")

    h("Phone and ID normalization", level=2)
    bullets([
        "Strip non-digits from phones; store a helper column for matching.",
        "Store account numbers as text in Sheets (apostrophe prefix in Excel export instructions).",
        "Trim spaces on names and addresses; optional standard abbreviations (St, Ave) if client agrees.",
    ])

    h("When to use Zapier vs Apps Script", level=2)
    p("Use Zapier for: new call row in Call_Log, simple field mapping, optional lookup by phone.")
    p("Use Apps Script for: full import upsert, duplicate logic, multi-match rules, menu tools, locking, and combining call + customer + note in one step.")

    h("If the client asks the eight screening questions", level=2)
    p("Answer honestly in the proposal; this is how you implement what you promise:")
    bullets([
        "RingCentral + Zapier + Sheets — Yes: call trigger → Call_Log → phone lookup → customer link.",
        "Sheets CRM — Yes: audit-first, preserve structure, improve matching and workflows.",
        "Excel import — Yes: staging → validate → upsert → log.",
        "Duplicates — Canonical keys, 2+ matches to review queue, no silent merge.",
        "Incoming calls — Caller ID normalized → lookup → unmatched queue for agents.",
        "Preserve scripts/tabs — Yes, with backup and written approval for removals.",
        "Milestones — Discovery → Mapping → Import → Duplicates → RingCentral → Docs/Test → Sign-off.",
        "Timeline — About 3–4 weeks for a typical setup; longer if Excel or scripts are messy.",
    ])

    h("Milestones and client sign-off", level=2)
    bullets([
        "M1 Discovery — Audit + backup + scope confirmed.",
        "M2 Mapping — Excel→CRM mapping approved.",
        "M3 Import — Staging + Script + successful test batch.",
        "M4 Duplicates — Review tab + scan working.",
        "M5 RingCentral — Zaps live + test calls verified.",
        "M6 Go-live — Agent/admin docs + training walkthrough.",
        "M7 Acceptance — Test_Log complete + Phase 1 sign-off.",
    ])

    h("Common problems and fixes", level=2)
    bullets([
        "Excel leading zeros lost — instruct text export; validate in staging.",
        "Zapier task limit — fewer Zaps, filter to answered/missed only, or Script webhook.",
        "Old Script conflicts — read triggers; namespace new functions; test on sandbox.",
        "Wrong caller ID — agent manually links call; track unmatched rate.",
        "Client changes columns mid-project — freeze mapping; change request for v1.1.",
    ])

    h("What not to do in Phase 1")
    bullets([
        "Rebuild the entire CRM from scratch.",
        "Remove tabs or scripts without written approval.",
        "Add Smith.ai, Google Calendar, or SMS automations.",
        "Auto-merge fuzzy name matches without human review.",
        "Go live without test calls and sample imports.",
    ])

    h("Checklist — job post deliverables")
    steps([
        "Reviewed Google Sheets CRM (audit document).",
        "Reviewed ISP Excel (mapping document).",
        "Import/update process live and documented.",
        "Duplicate prevention live and tested.",
        "RingCentral connected to Sheets via Zapier (and Script if needed).",
        "Call and retention status tracking with validated lists.",
        "Agent workflow documented and walked through.",
        "Tested with sample customers (Test_Log).",
        "Simple written instructions for agents and admin.",
    ])

    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build()
